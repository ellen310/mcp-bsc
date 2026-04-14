"""Tests for xlsx_handler.py."""

import os
import pytest
from openpyxl import Workbook

from mcp_bsc.xlsx_handler import parse_org_performance_sheet, create_personal_performance_sheet


@pytest.fixture()
def org_xlsx(tmp_path):
    """Create a minimal org performance sheet xlsx."""
    wb = Workbook()
    ws = wb.active
    # Header row
    ws.append(["카테고리", "성과항목", "지표", "목표치", "가중치"])
    # Data rows
    ws.append(["업무성과", "프로젝트 완료율", "완료 프로젝트 수", "90% 이상", "30%"])
    ws.append(["역량개발", "교육 이수", "교육 시간", "40시간 이상", "15%"])
    path = str(tmp_path / "org.xlsx")
    wb.save(path)
    return path


@pytest.fixture()
def personal_template(tmp_path):
    """Create a minimal personal performance template xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.append(["카테고리", "성과항목", "지표", "목표치", "달성내용", "근거자료", "점수"])
    path = str(tmp_path / "personal_template.xlsx")
    wb.save(path)
    return path


class TestParseOrgPerformanceSheet:
    def test_returns_list_of_dicts(self, org_xlsx):
        items = parse_org_performance_sheet(org_xlsx)
        assert isinstance(items, list)
        assert len(items) == 2

    def test_dict_has_expected_keys(self, org_xlsx):
        items = parse_org_performance_sheet(org_xlsx)
        for item in items:
            assert "category" in item
            assert "item" in item
            assert "indicator" in item
            assert "target" in item
            assert "weight" in item

    def test_parses_values_correctly(self, org_xlsx):
        items = parse_org_performance_sheet(org_xlsx)
        assert items[0]["category"] == "업무성과"
        assert items[0]["item"] == "프로젝트 완료율"
        assert items[1]["weight"] == "15%"

    def test_skips_header_row(self, org_xlsx):
        items = parse_org_performance_sheet(org_xlsx)
        categories = [i["category"] for i in items]
        assert "카테고리" not in categories


class TestCreatePersonalPerformanceSheet:
    def test_creates_file(self, personal_template, tmp_path):
        output = str(tmp_path / "output" / "personal.xlsx")
        data = [
            {
                "category": "업무성과",
                "item": "프로젝트",
                "indicator": "완료율",
                "target": "90%",
                "achievement": "92% 달성",
                "evidence": "프로젝트 보고서",
                "score": "A",
            }
        ]
        saved = create_personal_performance_sheet(personal_template, output, data)
        assert os.path.isfile(saved)

    def test_written_data_is_correct(self, personal_template, tmp_path):
        from openpyxl import load_workbook

        output = str(tmp_path / "personal_out.xlsx")
        data = [
            {
                "category": "역량개발",
                "item": "교육 이수",
                "indicator": "교육 시간",
                "target": "40시간",
                "achievement": "45시간 이수",
                "evidence": "교육 수료증",
                "score": "S",
            }
        ]
        create_personal_performance_sheet(personal_template, output, data)
        wb = load_workbook(output)
        ws = wb.active
        # Row 2 is the first data row
        assert ws.cell(row=2, column=1).value == "역량개발"
        assert ws.cell(row=2, column=5).value == "45시간 이수"
        assert ws.cell(row=2, column=7).value == "S"
        wb.close()
