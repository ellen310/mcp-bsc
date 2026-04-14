"""Generate sample xlsx templates for org and personal performance sheets."""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "templates")


def create_org_template() -> str:
    """Create org_performance_template.xlsx with headers and sample data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "조직성과관리표"

    headers = ["카테고리", "성과항목", "지표", "목표치", "가중치"]
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sample_data = [
        ["업무성과", "프로젝트 관리", "프로젝트 완료율", "90% 이상", "30%"],
        ["업무성과", "고객 만족도", "고객 만족도 점수", "4.5점 이상 (5점 만점)", "20%"],
        ["역량개발", "교육 이수", "연간 교육 시간", "40시간 이상", "15%"],
    ]

    for row_idx, row_data in enumerate(sample_data, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-fit column widths (approximate)
    column_widths = [15, 20, 25, 25, 10]
    for col, width in enumerate(column_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    output_path = os.path.join(TEMPLATES_DIR, "org_performance_template.xlsx")
    wb.save(output_path)
    print(f"Created: {output_path}")
    return output_path


def create_personal_template() -> str:
    """Create personal_performance_template.xlsx with headers only."""
    wb = Workbook()
    ws = wb.active
    ws.title = "개인성과관리표"

    headers = ["카테고리", "성과항목", "지표", "목표치", "달성내용", "근거자료", "점수"]
    header_font_white = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    column_widths = [15, 20, 25, 25, 40, 30, 10]
    for col, width in enumerate(column_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    output_path = os.path.join(TEMPLATES_DIR, "personal_performance_template.xlsx")
    wb.save(output_path)
    print(f"Created: {output_path}")
    return output_path


if __name__ == "__main__":
    create_org_template()
    create_personal_template()
    print("All templates created successfully.")
